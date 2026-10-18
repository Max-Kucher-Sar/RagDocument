import json
from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
import xlwt
from qdrant_client import QdrantClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.config import Settings
from app.documents import Unit, chunk_units, extract
from app.main import create_app
from app.ollama import OllamaClient
from app.service import GroundedAnswer, RagService, ServiceError, Statement

pytestmark = pytest.mark.filterwarnings('ignore:Payload indexes have no effect in the local Qdrant')


class FakeAI:
    def __init__(self):
        self.answer = GroundedAnswer(insufficient_evidence=False, statements=[
            Statement(text='Срок возврата — 14 дней.', source_ids=[1])])
        self.context = None

    def embed(self, texts, *, query=False):
        # Deterministic vectors: exercises real Qdrant storage and cosine search.
        return [[1.0, 0.0, 0.0, 0.0] for _ in texts]

    def generate(self, instructions, question, sources, response_model):
        self.context = dict(question=question, sources=sources)
        return self.answer

    def ensure_models(self, *, chat=False):
        pass

    def status(self):
        return dict(ollama_available=True, missing_models=[])

    def close(self):
        pass


@pytest.fixture
def setup(tmp_path):
    settings = Settings(_env_file=None, data_dir=tmp_path / 'documents',
                        embedding_dimensions=4, app_token='')
    ai = FakeAI()
    service = RagService(settings, qdrant=QdrantClient(path=str(tmp_path / 'qdrant')), llm=ai)
    yield settings, service, ai
    service.close()


def upload(service, text='Срок возврата товара — 14 дней. Условия возврата указаны в договоре.'):
    return service.ingest('policy.txt', text.encode())


def test_full_upload_ask_source_download_delete(setup):
    settings, service, ai = setup
    with TestClient(create_app(settings, service)) as client:
        assert client.get('/').status_code == 200
        assert client.get('/api/status').json()['provider'] == 'ollama'
        response = client.post('/api/documents', files={'file':('policy.txt', 'Срок возврата — 14 дней.'.encode())})
        assert response.status_code == 200, response.text
        doc = response.json()
        assert doc['chunks'] == 1
        answer = client.post('/api/ask', json={'question':'Какой срок возврата?', 'document_ids':[doc['id']]}).json()
        assert '[1]' in answer['answer']
        source = answer['sources'][0]
        assert source['document_id'] == doc['id']
        assert '14 дней' in source['quote']
        assert source['url'].startswith('/viewer.html?')
        assert client.get(source['url']).status_code == 200
        body = client.get(f"/api/documents/{doc['id']}/chunks/0").json()
        chunk = body['chunk']
        assert body['unit']['text'][chunk['start']:chunk['end']] == source['quote']
        assert '14 дней' in client.get(body['original_url']).content.decode()
        assert client.delete(f"/api/documents/{doc['id']}").status_code == 200
        assert client.get('/api/documents').json() == []
        assert service._db().count(settings.qdrant_collection).count == 0
        assert client.get(body['original_url']).status_code == 404


def test_filter_and_duplicate(setup):
    _, service, ai = setup
    first = upload(service)
    duplicate = upload(service)
    second = upload(service, 'Доставка выполняется за три дня.')
    assert duplicate['duplicate'] and duplicate['id'] == first['id']
    assert len(service.list_documents()) == 2
    result = service.ask('Как доставляют?', [second['id']])
    assert {source['document_id'] for source in result['sources']} == {second['id']}
    assert len(ai.context['sources']) == 1
    assert service.ask('Вопрос', [])['sources'] == []


def test_no_evidence_and_invalid_citations(setup):
    _, service, ai = setup
    upload(service)
    ai.answer = GroundedAnswer(insufficient_evidence=True, statements=[])
    assert service.ask('Неизвестный факт')['insufficient_evidence']
    for refs in [[], [99], [0], [-1]]:
        ai.answer = GroundedAnswer(insufficient_evidence=False, statements=[Statement(text='Факт', source_ids=refs)])
        with pytest.raises(ServiceError, match='некорректными ссылками'):
            service.ask('Вопрос')
    service.settings.score_threshold = 1.0
    # A threshold above cosine score is not available; use an orthogonal query vector.
    ai.embed = lambda *args, **kw: [[0., 1., 0., 0.]]
    assert service.ask('Другой вопрос')['insufficient_evidence']


def test_persistence(setup):
    settings, service, ai = setup
    doc = upload(service)
    service.close()
    service.qdrant = None
    reopened = RagService(settings, qdrant=QdrantClient(path=str(settings.data_dir.parent / 'qdrant')), llm=FakeAI())
    try:
        assert reopened.list_documents()[0]['id'] == doc['id']
        assert reopened.ask('Срок?')['sources'][0]['document_id'] == doc['id']
    finally:
        reopened.close()


def test_failed_batch_rolls_back(setup):
    settings, service, ai = setup
    settings.chunk_size = 200
    settings.chunk_overlap = 10
    original = ai.embed
    calls = 0

    def fail_second(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError('simulated API outage')
        return original(*args, **kwargs)

    ai.embed = fail_second
    with pytest.raises(RuntimeError, match='simulated'):
        upload(service, 'Договор и условия возврата. ' * 400)
    assert calls == 2
    assert service.list_documents() == []
    assert service._db().count(settings.qdrant_collection).count == 0
    assert not list(settings.data_dir.iterdir())


def test_auth_and_missing_model(setup):
    settings, _, _ = setup
    settings.app_token = 'test-secret'
    def respond(request):
        return httpx.Response(200, json={'models':[{'name':settings.chat_model}]})
    ollama = OllamaClient(settings, httpx.Client(base_url='http://ollama/', transport=httpx.MockTransport(respond)))
    service = RagService(settings, llm=ollama)
    with TestClient(create_app(settings, service)) as client:
        assert client.get('/api/documents').status_code == 401
        assert client.get('/api/documents/id/file').status_code == 401
        assert client.get('/api/documents/id/chunks/0').status_code == 401
        headers = {'Authorization':'Bearer test-secret'}
        assert client.get('/api/status', headers=headers).status_code == 200
        response = client.post('/api/documents', headers=headers, files={'file':('x.txt', b'text')})
        assert response.status_code == 503
        assert 'docker-compose run --rm ollama-init' in response.json()['detail']
        assert not list(settings.data_dir.iterdir())


def test_validation_and_model_mismatch(setup):
    settings, service, _ = setup
    upload(service)
    settings.embedding_dimensions = 8
    with pytest.raises(ServiceError, match='Размерность'):
        service.ask('Вопрос')
    settings.embedding_dimensions = 4
    settings.embedding_model = 'another-model'
    with pytest.raises(ServiceError, match='другая модель'):
        service.ask('Вопрос')
    with pytest.raises(ServiceError) as error:
        service.get_source('../secret', 0)
    assert error.value.status == 404
    with TestClient(create_app(settings, service)) as client:
        assert client.post('/api/ask', json={'question':'  '}).status_code == 400
        assert client.post('/api/ask', json={'question':'x'*4001}).status_code == 422


def test_chunk_coverage_and_coordinates():
    text = ('Строка с текстом для проверки.\n' * 150).strip()
    chunks = chunk_units([Unit(text, 'Текст')], size=200, overlap=30)
    covered = set()
    for chunk in chunks:
        assert chunk['text'] == text[chunk['start']:chunk['end']]
        assert len(chunk['text']) <= 200
        assert chunk['location'].startswith('Строки ')
        covered.update(range(chunk['start'],chunk['end']))
    assert covered == set(range(len(text)))


def test_docx_paragraph_and_table_locations():
    doc = Document()
    doc.add_paragraph('Первый абзац')
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Ячейка'
    doc.add_paragraph('Второй абзац')
    stream = BytesIO(); doc.save(stream)
    units = extract('file.docx', stream.getvalue())
    assert [unit.location for unit in units] == ['Абзац 1', 'Таблица 1, строка 1', 'Абзац 2']
    assert 'Ячейка' in units[1].text


def test_pdf_page_locations_and_scans():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'), NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 20 200 Td (Return policy: 14 days.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = BytesIO(); writer.write(output)
    units = extract('file.pdf', output.getvalue())
    assert len(units) == 1 and units[0].page == 2
    assert units[0].location == 'Страница 2'
    assert '14 days' in units[0].text
    blank = PdfWriter(); blank.add_blank_page(width=100, height=100)
    output = BytesIO(); blank.write(output)
    with pytest.raises(ValueError, match='OCR'):
        extract('scan.pdf', output.getvalue())


@pytest.mark.parametrize('filename,data', [('empty.txt',b''), ('bad.exe',b'abc'), ('bad.txt',b'\xff')])
def test_invalid_documents(filename, data):
    with pytest.raises(ValueError):
        extract(filename, data)


def test_native_ollama_protocol(setup):
    """Exercise HTTP serialization, JSON schema and parsing without real inference."""
    settings, service, _ = setup
    doc = upload(service)
    calls = []

    def respond(request):
        calls.append(request.url.path)
        if request.method == 'GET':
            return httpx.Response(200,json={'models':[{'name':settings.chat_model},{'name':settings.embedding_model}]})
        body = json.loads(request.content)
        if request.url.path == '/api/embed':
            assert body['truncate'] is False
            assert body['input'][0].startswith('Instruct:')
            return httpx.Response(200, json={'embeddings':[[1.,0.,0.,0.] for _ in body['input']]})
        assert request.url.path == '/api/chat'
        assert body['model'] == 'qwen3:4b-instruct'
        assert body['format']['type'] == 'object'
        assert body['stream'] is False and body['think'] is False
        answer = {'insufficient_evidence':False,'statements':[{'text':'Срок возврата — 14 дней.','source_ids':[1]}]}
        return httpx.Response(200,json={'message':{'content':json.dumps(answer)},'done':True,'done_reason':'stop'})

    service.llm = OllamaClient(settings, httpx.Client(base_url='http://ollama/',transport=httpx.MockTransport(respond)))
    result = service.ask('Какой срок возврата?', [doc['id']])
    assert result['sources'][0]['document_id'] == doc['id']
    assert result['answer'].endswith('[1]')
    assert calls == ['/api/tags', '/api/embed', '/api/chat']


def test_xlsx_coordinates_headers_and_formulas(setup):
    _, service, _ = setup
    book = Workbook(); sheet = book.active; sheet.title = 'Возвраты'
    sheet.append(['Товар', 'Срок', 'Остаток', 'Доступен', 'Расчёт'])
    sheet.append(['Чайник', '14 дней', 0, False, '=1+2'])
    book.create_sheet('Доставка').append(['Отправка', '3 дня'])
    buffer = BytesIO(); book.save(buffer)
    content = buffer.getvalue()
    units = extract('data.xlsx', content)
    assert len(units) == 3
    unit = units[1]
    assert unit.sheet == 'Возвраты' and unit.cell_range == 'A2:E2'
    assert 'B2 [первая заполненная строка, B1: Срок]: 14 дней' in unit.text
    assert 'C2 [первая заполненная строка, C1: Остаток]: 0' in unit.text
    assert 'FALSE' in unit.text
    assert '=1+2 [нет сохранённого результата' in unit.text
    doc = service.ingest('data.xlsx', content)
    source = service.get_source(doc['id'], 1)
    assert source['chunk']['sheet'] == 'Возвраты'
    assert source['chunk']['cell_range'] == 'A2:E2'
    assert 'строка 2' in source['chunk']['location']


def test_xls_coordinates_and_types():
    from datetime import datetime
    book = xlwt.Workbook(); sheet = book.add_sheet('Склад')
    for col, name in enumerate(['Товар', 'Цена', 'Остаток', 'Доступен', 'Дата']):
        sheet.write(0, col, name)
    sheet.write(2, 0, 'Чайник'); sheet.write(2, 1, 1250); sheet.write(2, 2, 0)
    sheet.write(2, 3, False)
    sheet.write(2, 4, datetime(2026, 10, 18), xlwt.easyxf(num_format_str='YYYY-MM-DD'))
    buffer = BytesIO(); book.save(buffer)
    units = extract('legacy.xls', buffer.getvalue())
    assert len(units) == 2
    assert units[1].sheet == 'Склад' and units[1].cell_range == 'A3:E3'
    assert '1250' in units[1].text and 'FALSE' in units[1].text
    assert '2026-10-18' in units[1].text


@pytest.mark.parametrize('code,message', [(404,'ollama-init'),(500,'Ollama отклонила')])
def test_ollama_http_errors(setup, code, message):
    settings, _, _ = setup
    with httpx.Client(base_url='http://ollama/',transport=httpx.MockTransport(
            lambda req: httpx.Response(code, json={'error':'test'}))) as http:
        client = OllamaClient(settings, http)
        with pytest.raises(ServiceError, match=message):
            client.embed(['текст'])


def test_ollama_invalid_and_truncated_answer(setup):
    settings, _, _ = setup
    for payload in [{'message':{'content':'not JSON'}}, {'done_reason':'length'}]:
        with httpx.Client(base_url='http://ollama/',transport=httpx.MockTransport(
                lambda req: httpx.Response(200, json=payload))) as http:
            with pytest.raises(ServiceError):
                OllamaClient(settings, http).generate('instructions','question',[],GroundedAnswer)


def test_ollama_unavailable(setup):
    settings, _, _ = setup
    def disconnected(request):
        raise httpx.ConnectError('not running', request=request)
    with httpx.Client(base_url='http://ollama/',transport=httpx.MockTransport(disconnected)) as http:
        client = OllamaClient(settings, http)
        assert not client.status()['ollama_available']
        with pytest.raises(ServiceError, match='Нет соединения'):
            client.embed(['text'])
