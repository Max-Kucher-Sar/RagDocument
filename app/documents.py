"""Extraction and chunking with stable, verifiable source coordinates."""
from dataclasses import asdict, dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import xlrd
from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

SUPPORTED = {'.pdf', '.docx', '.txt', '.md', '.xlsx', '.xls'}
MAX_TEXT = 2_000_000
MAX_CHUNKS = 5000
MAX_SHEET_ROWS = 100_000
MAX_SHEET_COLUMNS = 256
MAX_WORKBOOK_CELLS = 1_000_000


@dataclass
class Unit:
    text: str
    location: str
    page: int | None = None
    sheet: str | None = None
    cell_range: str | None = None


def _check_zip(content: bytes):
    with ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100_000_000:
            raise ValueError('Распакованный документ слишком большой.')


def _cell_text(value):
    if value is None or value == '':
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def _excel_units(sheets):
    """One row per unit, with actual cell addresses and explicit header provenance."""
    units, total, cells = [], 0, 0
    for sheet_name, rows in sheets:
        header = None
        header_row = None
        for row_no, values in enumerate(rows, 1):
            cells += len(values)
            if row_no > MAX_SHEET_ROWS or len(values) > MAX_SHEET_COLUMNS or cells > MAX_WORKBOOK_CELLS:
                raise ValueError('Excel слишком большой: максимум 100 000 строк, 256 столбцов '
                                 'на лист и 1 млн просматриваемых ячеек на книгу.')
            nonempty = [(col, _cell_text(value)) for col, value in enumerate(values, 1)
                        if value is not None and value != '']
            nonempty = [(col, text) for col, text in nonempty if text]
            if not nonempty:
                continue
            if header is None:
                header, header_row = dict(nonempty), row_no
            parts = []
            for col, value in nonempty:
                letter = get_column_letter(col)
                label = f'{letter}{row_no}'
                if row_no != header_row and col in header:
                    label += f' [первая заполненная строка, {letter}{header_row}: {header[col]}]'
                parts.append(f'{label}: {value}')
            cell_range = f'{get_column_letter(nonempty[0][0])}{row_no}:{get_column_letter(nonempty[-1][0])}{row_no}'
            text = f'Лист «{sheet_name}», строка {row_no}\n' + '\n'.join(parts)
            total += len(text)
            if total > MAX_TEXT:
                raise ValueError('Слишком много текста. Разделите документ на части.')
            units.append(Unit(text, f'Лист «{sheet_name}», строка {row_no}, {cell_range}',
                              sheet=sheet_name, cell_range=cell_range))
    return units


def _extract_xlsx(content):
    _check_zip(content)
    values_book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        formula_book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            def sheets():
                for values_sheet, formula_sheet in zip(values_book.worksheets, formula_book.worksheets):
                    if (values_sheet.max_row or 0) > MAX_SHEET_ROWS or (values_sheet.max_column or 0) > MAX_SHEET_COLUMNS:
                        raise ValueError('Excel превышает лимит строк или столбцов. Уберите лишнее форматирование или разделите лист.')

                    def rows(vs=values_sheet, fs=formula_sheet):
                        for cached, formulas in zip(vs.iter_rows(values_only=True), fs.iter_rows()):
                            result = []
                            for value, cell in zip(cached, formulas):
                                if value is None and cell.data_type == 'f':
                                    value = f'{cell.value} [нет сохранённого результата; формула не вычислялась]'
                                result.append(value)
                            yield result

                    yield values_sheet.title, rows()
            return _excel_units(sheets())
        finally:
            formula_book.close()
    finally:
        values_book.close()


def _extract_xls(content):
    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        def sheets():
            for sheet in book.sheets():
                if sheet.nrows > MAX_SHEET_ROWS or sheet.ncols > MAX_SHEET_COLUMNS:
                    raise ValueError('Excel превышает лимит строк или столбцов.')

                def rows(ws=sheet):
                    for row in ws.get_rows():
                        values = []
                        for cell in row:
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                value = bool(cell.value)
                            elif cell.ctype == xlrd.XL_CELL_ERROR:
                                value = xlrd.error_text_from_code.get(cell.value, '#ERROR')
                            elif cell.ctype == xlrd.XL_CELL_NUMBER and cell.value.is_integer():
                                value = int(cell.value)
                            else:
                                value = cell.value
                            values.append(value)
                        yield values

                yield sheet.name, rows()
        return _excel_units(sheets())
    finally:
        book.release_resources()


def extract(filename: str, content: bytes) -> list[Unit]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError('Поддерживаются PDF, DOCX, TXT, MD, XLSX и XLS.')
    if suffix == '.pdf':
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('PDF защищён паролем. Загрузите открытый файл.')
        if len(reader.pages) > 2000:
            raise ValueError('Максимум 2000 страниц в одном PDF.')
        units = []
        total = 0
        for i, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or '').strip()
            total += len(text)
            if total > MAX_TEXT:
                raise ValueError('Слишком много текста. Разделите документ на части.')
            if text:
                units.append(Unit(text, f'Страница {i}', i))
    elif suffix == '.docx':
        # Reject oversized ZIP contents before python-docx decompresses them.
        _check_zip(content)
        document = Document(BytesIO(content))
        units = []
        paragraph = table = 0
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                table += 1
                for row_no, row in enumerate(block.rows, 1):
                    text = ' | '.join(cell.text for cell in row.cells).strip()
                    if text:
                        units.append(Unit(text, f'Таблица {table}, строка {row_no}'))
            else:
                paragraph += 1
                if block.text.strip():
                    units.append(Unit(block.text.strip(), f'Абзац {paragraph}'))
    elif suffix == '.xlsx':
        units = _extract_xlsx(content)
    elif suffix == '.xls':
        units = _extract_xls(content)
    else:
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('TXT/MD должен быть в кодировке UTF-8.') from exc
        text = text.replace('\r\n', '\n').replace('\r', '\n').strip()
        units = [Unit(text, 'Текст')] if text else []
    if sum(len(unit.text) for unit in units) > MAX_TEXT:
        raise ValueError('Слишком много текста. Разделите документ на части.')
    if not units:
        raise ValueError('В документе нет текста. Для сканов PDF сначала выполните OCR.')
    return units


def chunk_units(units: list[Unit], size: int, overlap: int) -> list[dict]:
    if size <= 0 or not 0 <= overlap < size:
        raise ValueError('Некорректный размер фрагмента или перекрытие.')
    chunks = []
    for unit_index, unit in enumerate(units):
        start = 0
        while start < len(unit.text):
            end = min(start + size, len(unit.text))
            if end < len(unit.text):
                boundary = max(unit.text.rfind('\n', start, end), unit.text.rfind(' ', start, end))
                if boundary > start + size // 2:
                    end = boundary
            text = unit.text[start:end]
            if text.strip():
                location = unit.location
                if location == 'Текст':
                    first = unit.text.count('\n', 0, start) + 1
                    last = unit.text.count('\n', 0, end) + 1
                    location = f'Строки {first}–{last}'
                chunks.append(dict(index=len(chunks), unit_index=unit_index,
                                   start=start, end=end, text=text,
                                   location=location, page=unit.page,
                                   sheet=unit.sheet, cell_range=unit.cell_range))
                if len(chunks) > MAX_CHUNKS:
                    raise ValueError('Слишком много фрагментов. Разделите документ.')
            if end == len(unit.text):
                break
            start = max(start + 1, end - overlap)
    return chunks


def serialize_units(units: list[Unit]) -> list[dict]:
    return [asdict(unit) for unit in units]

